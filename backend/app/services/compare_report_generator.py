from __future__ import annotations

import tempfile
from datetime import datetime
from pathlib import Path

from fastapi.responses import FileResponse
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, PageBreak
from app.services.visual_report_theme import (
    header_footer as _premium_footer, hero as _premium_hero, kpi_cards as _premium_kpis,
    section_title as _premium_section, notice_box as _premium_notice, bar_list as _premium_bar_list,
    text_panel as _premium_text_panel, table_style as _premium_table_style, short as _p_short,
    brl as _p_brl, pct as _p_pct, styles as _premium_styles
)

RED = RGBColor(197, 38, 37)
DARK = RGBColor(31, 41, 55)
GRAY = RGBColor(105, 112, 120)
LIGHT = RGBColor(244, 246, 248)
WHITE = RGBColor(255, 255, 255)
MID = RGBColor(226, 230, 235)


def _brl(value) -> str:
    try:
        return f"R$ {float(value or 0):,.0f}".replace(",", ".")
    except Exception:
        return "R$ 0"


def _pct(value) -> str:
    try:
        return f"{float(value or 0):.1f}%"
    except Exception:
        return "0,0%"

def _scope_label(value: str | None) -> str:
    scope = (value or "market").strip().lower()
    return {
        "market": "Mercado amplo",
        "preserved": "Somente com evidência preservada",
        "advertising": "Somente publicidade classificada",
        "auditavel": "Somente checking auditável",
    }.get(scope, "Mercado amplo")


def _quality_notice(data: dict) -> str | None:
    notice = data.get("quality_notice")
    if notice:
        return str(notice)
    totals = data.get("totals", {}) or {}
    filters = data.get("filters", {}) or {}
    scope = str(filters.get("evidence_scope") or "market").lower()
    items = int(float(totals.get("items") or 0))
    auditables = int(float(totals.get("auditables") or 0))
    if items > 0 and auditables == 0 and scope != "auditavel":
        return (
            f"Leitura referencial: o recorte possui {items} evidência(s) de mercado, "
            "mas nenhuma evidência auditável para checking."
        )
    if scope == "auditavel" and items == 0:
        return "Nenhuma evidência auditável foi encontrada para checking no recorte selecionado."
    return None


def _short(text, size=42):
    value = str(text or "—")
    return value if len(value) <= size else value[: size - 1] + "…"


def _report_name(prefix: str, project_id: str, ext: str) -> str:
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    safe_project = str(project_id).replace("/", "-")[:36]
    return f"{prefix}_{safe_project}_{stamp}.{ext}"


def build_compare_pdf(data: dict, output_path: str) -> str:
    """Premium Intel Comparativo PDF aligned with Editorial V20 visual system."""
    doc = SimpleDocTemplate(
        output_path,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=10 * mm,
        bottomMargin=14 * mm,
    )

    filters = data.get("filters", {}) or {}
    totals = data.get("totals", {}) or {}
    advertisers = data.get("advertisers", []) or []
    pairs = data.get("pairs", []) or []
    overlaps = data.get("portal_overlap", []) or []
    insights = data.get("insights", []) or []
    scope_label = filters.get("evidence_scope_label") or _scope_label(filters.get("evidence_scope"))
    notice = _quality_notice(data)
    segment = filters.get("segment_name") or "Todos os segmentos"

    story: list = []
    story.append(_premium_hero(
        "TV Fiscal Intel Comparativo",
        "Cruzamento competitivo por segmento, anunciante, portal, formato, evidência e investimento estimado",
        project_label=f"Projeto {data.get('project_id', '—')} · Segmento: {segment}",
        kicker="INTEL COMPARATIVO · TV FISCAL WEBMONITOR",
    ))
    story.append(Spacer(1, 4 * mm))
    story.append(_premium_kpis([
        ("Anunciantes", totals.get("advertisers", 0), "comparados"),
        ("Evidências", totals.get("items", 0), "recorte"),
        ("Auditáveis", totals.get("auditables", 0), "checking"),
        ("Portais", totals.get("portals", 0), "capilaridade"),
        ("Investimento", _p_brl(totals.get("investment", 0)), "estimado"),
        ("Modo", scope_label, "análise"),
    ]))
    story.append(Spacer(1, 4 * mm))
    if notice:
        story.append(_premium_notice(notice, tone="warning"))
        story.append(Spacer(1, 3 * mm))

    story.append(Table([[
        _premium_bar_list("Ranking por anunciante", advertisers, "advertiser", "items", sub_key="investment_label", limit=6),
        _premium_bar_list("Share por investimento", advertisers, "advertiser", "share_investment_percent", sub_key="share_label", max_value=100, limit=6),
        _premium_text_panel("Insights competitivos", insights[:7] or ["Sem insights disponíveis para o recorte atual."], width=64 * mm),
    ]], colWidths=[94 * mm, 94 * mm, 66 * mm], rowHeights=[72 * mm], style=TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ])))

    story.append(PageBreak())
    story.append(_premium_section("Ranking comparativo por anunciante"))
    adv_rows = [["#", "Anunciante", "Segmento", "Evid.", "Audit.", "Invest.", "Share inv.", "Portais", "Força"]]
    for index, item in enumerate(advertisers[:14], start=1):
        inv = item.get("investment", 0)
        item.setdefault("investment_label", _p_brl(inv))
        item.setdefault("share_label", _p_pct(item.get("share_investment_percent", 0)))
        adv_rows.append([
            index,
            _p_short(item.get("advertiser"), 28),
            _p_short(item.get("segment_name"), 20),
            item.get("items", 0),
            item.get("auditables", 0),
            _p_brl(inv),
            _p_pct(item.get("share_investment_percent", 0)),
            len(item.get("portals", []) or []),
            item.get("strength_score", 0),
        ])
    if len(adv_rows) == 1:
        adv_rows.append(["—", "Sem dados", "—", 0, 0, _p_brl(0), _p_pct(0), 0, 0])
    table = Table(adv_rows, colWidths=[10 * mm, 55 * mm, 40 * mm, 18 * mm, 18 * mm, 30 * mm, 24 * mm, 20 * mm, 20 * mm])
    table.setStyle(_premium_table_style())
    story.append(table)

    story.append(_premium_section("Matriz de confronto direto"))
    pair_rows = [["Confronto", "Líder", "Equilíbrio", "Classificação", "Delta inv.", "Delta evid."]]
    for item in pairs[:12]:
        pair_rows.append([
            f"{_p_short(item.get('left'), 18)} x {_p_short(item.get('right'), 18)}",
            _p_short(item.get("leader"), 22),
            item.get("balance_score", 0),
            item.get("relation", "—"),
            _p_brl(item.get("investment_delta", 0)),
            item.get("items_delta", 0),
        ])
    if len(pair_rows) == 1:
        pair_rows.append(["Sem confronto competitivo válido", "—", "—", "—", "—", "—"])
    table = Table(pair_rows, colWidths=[70 * mm, 45 * mm, 28 * mm, 44 * mm, 30 * mm, 25 * mm])
    table.setStyle(_premium_table_style())
    story.append(table)

    story.append(_premium_section("Sobreposição de portais"))
    overlap_rows = [["Confronto", "Portais em comum", "Qtd."]]
    for item in overlaps[:12]:
        overlap_rows.append([
            f"{_p_short(item.get('left'), 24)} x {_p_short(item.get('right'), 24)}",
            _p_short(", ".join(item.get("common_portals", []) or []), 100),
            item.get("count", 0),
        ])
    if len(overlap_rows) == 1:
        overlap_rows.append(["Sem sobreposição", "—", "0"])
    table = Table(overlap_rows, colWidths=[80 * mm, 135 * mm, 20 * mm])
    table.setStyle(_premium_table_style())
    story.append(table)

    doc.build(story, onFirstPage=_premium_footer, onLaterPages=_premium_footer)
    return output_path

def _base_table_style() -> TableStyle:
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#b00020")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
        ("BACKGROUND", (0, 1), (-1, -1), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f9fafb")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ])


def _add_bg(slide, color=WHITE):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)


def _add_text(slide, left, top, width, height, text, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Arial"
    return box


def _add_title(slide, title, subtitle=None):
    _add_text(slide, Inches(0.65), Inches(0.62), Inches(12), Inches(0.45), title, 26, True, DARK)
    if subtitle:
        _add_text(slide, Inches(0.68), Inches(1.12), Inches(11.7), Inches(0.35), subtitle, 12, False, GRAY)


def _add_card(slide, left, top, width, height, label, value):
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = MID
    _add_text(slide, left + Inches(0.18), top + Inches(0.13), width - Inches(0.3), Inches(0.25), label, 9, True, GRAY)
    _add_text(slide, left + Inches(0.18), top + Inches(0.43), width - Inches(0.3), Inches(0.5), value, 20, True, RED)


def _add_table_slide(slide, title, headers, rows, left=Inches(0.6), top=Inches(1.65), width=Inches(12.1), height=Inches(4.8)):
    _add_text(slide, left, top - Inches(0.38), width, Inches(0.3), title, 15, True, RED)
    rows = rows[:8] if rows else [["—"] * len(headers)]
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), left, top, width, height)
    table = table_shape.table
    for c, header in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = str(header)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RED
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(8)
                r.font.bold = True
                r.font.color.rgb = WHITE
                r.font.name = "Arial"
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r_idx % 2 else LIGHT
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
                    r.font.color.rgb = DARK
                    r.font.name = "Arial"


def build_compare_pptx(data: dict, output_path: str) -> str:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    filters = data.get("filters", {}) or {}
    totals = data.get("totals", {}) or {}
    advertisers = data.get("advertisers", []) or []
    pairs = data.get("pairs", []) or []
    overlaps = data.get("portal_overlap", []) or []
    insights = data.get("insights", []) or []
    scope_label = filters.get("evidence_scope_label") or _scope_label(filters.get("evidence_scope"))
    notice = _quality_notice(data)

    # Capa
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, LIGHT)
    hero = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.35), Inches(12.0), Inches(2.15))
    hero.fill.solid()
    hero.fill.fore_color.rgb = RED
    hero.line.fill.background()
    _add_text(slide, Inches(1.05), Inches(1.82), Inches(10.8), Inches(0.42), "TV Fiscal WebMonitor", 18, True, WHITE)
    _add_text(slide, Inches(1.05), Inches(2.25), Inches(10.8), Inches(0.62), "Intel Comparativo", 32, True, WHITE)
    _add_text(slide, Inches(1.05), Inches(3.02), Inches(10.8), Inches(0.32), f"Segmento: {filters.get('segment_name') or 'Todos'}", 13, False, WHITE)
    _add_text(slide, Inches(0.9), Inches(4.25), Inches(11.8), Inches(0.36), f"Projeto: {data.get('project_id', '—')}", 13, False, DARK)
    _add_text(slide, Inches(0.9), Inches(4.62), Inches(11.8), Inches(0.36), f"Modo de análise: {scope_label}", 13, True, RED)
    _add_text(slide, Inches(0.9), Inches(5.0), Inches(11.8), Inches(0.36), f"Investimento estimado: {_brl(totals.get('investment'))} · Evidências: {totals.get('items', 0)} · Auditáveis: {totals.get('auditables', 0)}", 16, True, DARK)

    # Resumo
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, WHITE)
    _add_title(slide, "Resumo do confronto", "Indicadores executivos do recorte competitivo")
    _add_card(slide, Inches(0.7), Inches(1.75), Inches(2.35), Inches(1.15), "Anunciantes", str(totals.get("advertisers", 0)))
    _add_card(slide, Inches(3.25), Inches(1.75), Inches(2.35), Inches(1.15), "Evidências", str(totals.get("items", 0)))
    _add_card(slide, Inches(5.8), Inches(1.75), Inches(2.35), Inches(1.15), "Auditáveis", str(totals.get("auditables", 0)))
    _add_card(slide, Inches(8.35), Inches(1.75), Inches(2.35), Inches(1.15), "Portais", str(totals.get("portals", 0)))
    _add_card(slide, Inches(10.9), Inches(1.75), Inches(1.75), Inches(1.15), "Invest.", _brl(totals.get("investment", 0)))
    y = Inches(3.22)
    if notice:
        _add_text(slide, Inches(0.8), y, Inches(11.8), Inches(0.58), f"Atenção: {notice}", 12, True, RGBColor(154, 103, 0))
        y += Inches(0.68)
    _add_text(slide, Inches(0.8), y - Inches(0.25), Inches(11.8), Inches(0.3), "Insights automáticos", 16, True, RED)
    for item in (insights or ["Sem insights disponíveis."])[:6]:
        _add_text(slide, Inches(0.95), y, Inches(11.5), Inches(0.35), f"• {item}", 13, False, DARK)
        y += Inches(0.45)

    # Ranking
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, LIGHT)
    _add_title(slide, "Ranking por anunciante", "Share, investimento, evidências e força competitiva")
    rows = [
        [idx + 1, _short(a.get("advertiser"), 22), a.get("items", 0), a.get("auditables", 0), _brl(a.get("investment", 0)), _pct(a.get("share_investment_percent", 0)), a.get("strength_score", 0)]
        for idx, a in enumerate(advertisers[:8])
    ]
    _add_table_slide(slide, "Top anunciantes", ["#", "Anunciante", "Evid.", "Audit.", "Invest.", "Share", "Força"], rows)

    # Matriz
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, WHITE)
    _add_title(slide, "Matriz de confronto direto", "Disputas equilibradas, liderança e relação competitiva")
    rows = [
        [f"{_short(p.get('left'), 18)} x {_short(p.get('right'), 18)}", _short(p.get("leader"), 20), p.get("balance_score", 0), p.get("relation", "—"), _brl(p.get("investment_delta", 0)), p.get("items_delta", 0)]
        for p in pairs[:8]
    ]
    _add_table_slide(slide, "Confrontos", ["Confronto", "Líder", "Equil.", "Classificação", "Delta Inv.", "Delta Evid."], rows)

    # Portais
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, LIGHT)
    _add_title(slide, "Sobreposição de portais", "Onde os concorrentes aparecem nos mesmos canais")
    rows = [
        [f"{_short(o.get('left'), 20)} x {_short(o.get('right'), 20)}", _short(", ".join(o.get("common_portals", []) or []), 60), o.get("count", 0)]
        for o in overlaps[:8]
    ]
    _add_table_slide(slide, "Portais em comum", ["Confronto", "Portais", "Qtd."], rows)

    prs.save(output_path)
    return output_path


def build_compare_pdf_response(data: dict) -> FileResponse:
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    tmp.close()
    build_compare_pdf(data, tmp.name)
    project_id = data.get("project_id", "projeto")
    return FileResponse(tmp.name, media_type="application/pdf", filename=_report_name("intel_comparativo", project_id, "pdf"))


def build_compare_pptx_response(data: dict) -> FileResponse:
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pptx")
    tmp.close()
    build_compare_pptx(data, tmp.name)
    project_id = data.get("project_id", "projeto")
    return FileResponse(tmp.name, media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation", filename=_report_name("intel_comparativo", project_id, "pptx"))
