from __future__ import annotations

import csv
import io
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from typing import Any

from sqlalchemy.orm import Session
from sqlalchemy import and_

from app.db.models import NotificationLog
from app.services.notification_service import serialize_log, _sla_status, _sla_remaining_minutes
from app.utils.timeutils import display_local, utc_now


def _uuid(value: Any):
    import uuid
    if not value:
        return None
    if isinstance(value, uuid.UUID):
        return value
    return uuid.UUID(str(value))


def _parse_dt(value: str | None) -> datetime | None:
    if not value:
        return None
    text = str(value).strip()
    if not text:
        return None
    # Aceita YYYY-MM-DD ou ISO simples. Datas sem hora usam início/fim no chamador.
    try:
        if len(text) == 10:
            return datetime.fromisoformat(text)
        return datetime.fromisoformat(text.replace("Z", "+00:00")).replace(tzinfo=None)
    except Exception:
        return None


def _period_filter(query, start: datetime | None, end: datetime | None):
    if start:
        query = query.filter(NotificationLog.created_at >= start)
    if end:
        query = query.filter(NotificationLog.created_at < end)
    return query


def _status(row: NotificationLog) -> str:
    return str(getattr(row, "status", None) or "").lower()


def _alert_status(row: NotificationLog) -> str:
    return str(getattr(row, "alert_status", None) or "aberto").lower()


def _severity(row: NotificationLog) -> str:
    return str(getattr(row, "severity", None) or "informativo").lower()


def _channel(row: NotificationLog) -> str:
    return str(getattr(row, "channel", None) or "painel").lower()


def _category(row: NotificationLog) -> str:
    return str(getattr(row, "category", None) or getattr(row, "event_type", None) or "sem categoria")


def _avg_minutes(rows: list[NotificationLog], attr_start: str, attr_end: str) -> float:
    values: list[float] = []
    for r in rows:
        start = getattr(r, attr_start, None)
        end = getattr(r, attr_end, None)
        if start and end:
            try:
                values.append(max(0, (end - start).total_seconds() / 60.0))
            except Exception:
                pass
    return round(sum(values) / len(values), 1) if values else 0.0


def _count_list(counter: Counter, key_name: str, limit: int = 12) -> list[dict]:
    return [{key_name: k, "count": v} for k, v in counter.most_common(limit)]


def _day_key(dt: datetime | None) -> str:
    return display_local(dt).split(",")[0] if dt else "sem data"


def build_alerts_management_report(
    db: Session,
    project_id: str,
    start_date: str | None = None,
    end_date: str | None = None,
    limit: int = 5000,
) -> dict:
    project_uuid = _uuid(project_id)
    start = _parse_dt(start_date)
    end = _parse_dt(end_date)
    if end and len(str(end_date or "")) == 10:
        end = end + timedelta(days=1)

    query = db.query(NotificationLog).filter(NotificationLog.project_id == project_uuid)
    query = _period_filter(query, start, end)
    rows = query.order_by(NotificationLog.created_at.desc()).limit(limit).all()

    active = [r for r in rows if _alert_status(r) != "resolvido"]
    resolved = [r for r in rows if _alert_status(r) == "resolvido"]
    acknowledged = [r for r in rows if _alert_status(r) == "ciente"]
    snoozed = [r for r in rows if _alert_status(r) == "adiado"]
    overdue = [r for r in active if _sla_status(r) == "vencido"]
    due_soon = [r for r in active if _sla_status(r) == "vence_em_breve"]
    unassigned = [r for r in active if not getattr(r, "assigned_to", None)]
    escalated = [r for r in rows if (getattr(r, "escalation_count", 0) or 0) > 0]
    sent = [r for r in rows if _status(r) == "sent"]
    errors = [r for r in rows if _status(r) == "error"]
    dry = [r for r in rows if _status(r) == "dry_run"]
    panel = [r for r in rows if _status(r) == "registered" or _channel(r) == "painel"]

    owner_counter = Counter((getattr(r, "assigned_to", None) or "Sem responsável") for r in active)
    severity_counter = Counter(_severity(r) for r in rows)
    channel_counter = Counter(_channel(r) for r in rows)
    status_counter = Counter(_status(r) or "sem status" for r in rows)
    alert_status_counter = Counter(_alert_status(r) for r in rows)
    category_counter = Counter(_category(r) for r in rows)
    provider_counter = Counter((getattr(r, "provider", None) or "sem provider") for r in rows)
    day_counter = Counter(_day_key(getattr(r, "created_at", None)) for r in rows)

    # Reincidência simples por tema/marca: usa categoria + palavras relevantes no título.
    recurrent_terms = Counter()
    stop = {"tv", "fiscal", "alerta", "teste", "notificacao", "notificação", "de", "da", "do", "para", "por", "em", "e", "o", "a", "um", "uma", "com"}
    for r in rows:
        text = f"{getattr(r, 'category', '')} {getattr(r, 'title', '')}".lower()
        import re
        for token in re.findall(r"[a-zA-ZÀ-ÿ0-9]{4,}", text):
            if token not in stop:
                recurrent_terms[token] += 1

    avg_ack = _avg_minutes([r for r in rows if getattr(r, "acknowledged_at", None)], "created_at", "acknowledged_at")
    avg_resolve = _avg_minutes([r for r in rows if getattr(r, "resolved_at", None)], "created_at", "resolved_at")
    resolution_rate = round((len(resolved) / len(rows) * 100), 1) if rows else 0.0
    overdue_rate = round((len(overdue) / len(active) * 100), 1) if active else 0.0
    error_rate = round((len(errors) / len(rows) * 100), 1) if rows else 0.0

    critical_queue = sorted(
        active,
        key=lambda r: (
            0 if _sla_status(r) == "vencido" else 1,
            getattr(r, "sla_due_at", None) or utc_now(),
            -({"critico": 4, "alto": 3, "medio": 2}.get(_severity(r), 1)),
        ),
    )[:30]

    period_label = "Todo o histórico"
    if start_date or end_date:
        period_label = f"{start_date or 'início'} a {end_date or 'hoje'}"

    return {
        "project_id": project_id,
        "period": {"start_date": start_date, "end_date": end_date, "label": period_label},
        "generated_at": display_local(utc_now()),
        "summary": {
            "total": len(rows),
            "active": len(active),
            "resolved": len(resolved),
            "acknowledged": len(acknowledged),
            "snoozed": len(snoozed),
            "overdue": len(overdue),
            "due_soon": len(due_soon),
            "unassigned": len(unassigned),
            "escalated": len(escalated),
            "sent": len(sent),
            "errors": len(errors),
            "dry_run": len(dry),
            "panel": len(panel),
            "resolution_rate": resolution_rate,
            "overdue_rate": overdue_rate,
            "error_rate": error_rate,
            "avg_ack_minutes": avg_ack,
            "avg_resolution_minutes": avg_resolve,
        },
        "by_owner": _count_list(owner_counter, "owner"),
        "by_severity": _count_list(severity_counter, "severity"),
        "by_channel": _count_list(channel_counter, "channel"),
        "by_status": _count_list(status_counter, "status"),
        "by_alert_status": _count_list(alert_status_counter, "alert_status"),
        "by_category": _count_list(category_counter, "category"),
        "by_provider": _count_list(provider_counter, "provider"),
        "timeline": [{"date": k, "count": v} for k, v in sorted(day_counter.items())],
        "recurrent_terms": _count_list(recurrent_terms, "term"),
        "critical_queue": [serialize_log(r) for r in critical_queue],
        "latest": [serialize_log(r) for r in rows[:40]],
        "executive_reading": _executive_reading(len(rows), len(resolved), len(overdue), len(unassigned), len(errors), avg_resolve, resolution_rate, overdue_rate),
    }


def _executive_reading(total: int, resolved: int, overdue: int, unassigned: int, errors: int, avg_resolve: float, resolution_rate: float, overdue_rate: float) -> list[str]:
    if total == 0:
        return ["Não há alertas no período selecionado."]
    lines = [f"Foram analisados {total} alerta(s), com taxa de resolução de {resolution_rate:.1f}%."]
    if overdue:
        lines.append(f"Existem {overdue} alerta(s) vencido(s), equivalentes a {overdue_rate:.1f}% da fila ativa.")
    if unassigned:
        lines.append(f"Há {unassigned} alerta(s) sem responsável definido; recomenda-se atribuição imediata.")
    if errors:
        lines.append(f"Foram registrados {errors} erro(s) de envio/notificação; revisar provedores e logs.")
    if avg_resolve:
        lines.append(f"O tempo médio de resolução no período foi de {avg_resolve:.1f} minuto(s).")
    if not overdue and not errors:
        lines.append("A operação apresenta controle adequado de SLA no recorte atual.")
    return lines


def alerts_report_csv(report: dict) -> str:
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(["Relatório Gerencial de Alertas e SLA", report.get("project_id"), report.get("period", {}).get("label"), report.get("generated_at")])
    writer.writerow([])
    writer.writerow(["Indicador", "Valor"])
    for key, value in (report.get("summary") or {}).items():
        writer.writerow([key, value])
    writer.writerow([])
    writer.writerow(["Últimos alertas"])
    writer.writerow(["Data", "Canal", "Status", "Status alerta", "SLA", "Severidade", "Responsável", "Título", "Resposta"])
    for row in report.get("latest", []):
        writer.writerow([
            row.get("created_at_display") or row.get("created_at"), row.get("channel"), row.get("status"), row.get("alert_status"), row.get("sla_status"), row.get("severity"), row.get("assigned_to") or "Sem responsável", row.get("title"), row.get("provider_response"),
        ])
    return output.getvalue()


def generate_alerts_report_pdf(report: dict) -> bytes:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, PageBreak
    from app.services.visual_report_theme import hero, kpi_cards, section_title, table_style, styles, header_footer, safe, short

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=12*mm, rightMargin=12*mm, topMargin=10*mm, bottomMargin=15*mm)
    s = styles()
    story = []
    story.append(hero("Relatório Gerencial de Alertas e SLA", "Operação, produtividade, vencimentos, escalonamentos e canais de notificação", f"Projeto {report.get('project_id')} · {report.get('period', {}).get('label')}", "ALERT OPS · TV FISCAL WEBMONITOR"))
    story.append(Spacer(1, 5*mm))
    sm = report.get("summary", {})
    story.append(kpi_cards([
        ("Alertas", sm.get("total", 0), "período"),
        ("Ativos", sm.get("active", 0), "fila"),
        ("Resolvidos", sm.get("resolved", 0), f"{sm.get('resolution_rate', 0)}%"),
        ("Vencidos", sm.get("overdue", 0), f"{sm.get('overdue_rate', 0)}%"),
        ("Sem responsável", sm.get("unassigned", 0), "atribuir"),
        ("Erros", sm.get("errors", 0), f"{sm.get('error_rate', 0)}%"),
    ]))
    story.append(Spacer(1, 5*mm))
    story.append(section_title("Leitura executiva"))
    for line in report.get("executive_reading", []):
        story.append(Paragraph(f"• {safe(line)}", s["body"]))
    story.append(Spacer(1, 4*mm))

    # Tabelas compactas lado a lado
    def simple_table(title: str, rows: list[dict], key: str) -> Table:
        data = [[Paragraph(title, s["body"]), Paragraph("Qtd.", s["body"])]]
        for r in rows[:8]:
            data.append([Paragraph(short(r.get(key), 34), s["small"]), Paragraph(str(r.get("count", 0)), s["small"])])
        if len(data) == 1:
            data.append([Paragraph("Sem dados", s["small"]), "0"])
        t = Table(data, colWidths=[54*mm, 16*mm])
        t.setStyle(table_style())
        return t

    story.append(Table([[simple_table("Por responsável", report.get("by_owner", []), "owner"), simple_table("Por severidade", report.get("by_severity", []), "severity"), simple_table("Por canal", report.get("by_channel", []), "channel")]], colWidths=[82*mm,82*mm,82*mm]))
    story.append(Spacer(1, 5*mm))
    story.append(Table([[simple_table("Por status operacional", report.get("by_alert_status", []), "alert_status"), simple_table("Por categoria", report.get("by_category", []), "category"), simple_table("Reincidência termos/temas", report.get("recurrent_terms", []), "term")]], colWidths=[82*mm,82*mm,82*mm]))
    story.append(Spacer(1, 5*mm))
    story.append(section_title("Fila crítica de SLA"))
    data = [["Data", "SLA", "Severidade", "Responsável", "Título"]]
    for row in report.get("critical_queue", [])[:12]:
        data.append([
            row.get("created_at_display") or "—",
            row.get("sla_status") or "—",
            row.get("severity") or "—",
            row.get("assigned_to") or "Sem responsável",
            short(row.get("title"), 76),
        ])
    if len(data) == 1:
        data.append(["—", "—", "—", "—", "Nenhum alerta crítico no período."])
    t = Table(data, colWidths=[32*mm, 27*mm, 25*mm, 42*mm, 124*mm])
    t.setStyle(table_style())
    story.append(t)
    story.append(PageBreak())

    story.append(section_title("Últimos alertas registrados"))
    data = [["Data", "Canal", "Status", "SLA", "Severidade", "Título", "Resposta"]]
    for row in report.get("latest", [])[:24]:
        data.append([
            row.get("created_at_display") or "—",
            row.get("channel") or "—",
            row.get("status") or "—",
            row.get("sla_status") or "—",
            row.get("severity") or "—",
            short(row.get("title"), 46),
            short(row.get("provider_response"), 68),
        ])
    if len(data) == 1:
        data.append(["—", "—", "—", "—", "—", "Sem alertas", "—"])
    t = Table(data, colWidths=[30*mm, 20*mm, 22*mm, 24*mm, 24*mm, 60*mm, 70*mm])
    t.setStyle(table_style())
    story.append(t)
    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)
    return buffer.getvalue()


def generate_alerts_report_pptx(report: dict) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    navy = RGBColor(7, 26, 47)
    red = RGBColor(193, 18, 31)
    gray = RGBColor(100, 116, 139)

    def add_title(slide, title, subtitle=""):
        bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.78))
        bg.fill.solid(); bg.fill.fore_color.rgb = navy; bg.line.color.rgb = navy
        box = slide.shapes.add_textbox(Inches(0.45), Inches(0.16), Inches(11.8), Inches(0.5))
        p = box.text_frame.paragraphs[0]
        p.text = title; p.font.size = Pt(22); p.font.bold = True; p.font.color.rgb = RGBColor(255,255,255)
        if subtitle:
            sub = slide.shapes.add_textbox(Inches(0.47), Inches(0.82), Inches(12), Inches(0.28))
            q = sub.text_frame.paragraphs[0]; q.text = subtitle; q.font.size = Pt(10); q.font.color.rgb = gray

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Relatório Gerencial de Alertas e SLA", f"Projeto {report.get('project_id')} · {report.get('period', {}).get('label')}")
    sm = report.get("summary", {})
    metrics = [("Alertas", sm.get("total",0)), ("Ativos", sm.get("active",0)), ("Resolvidos", sm.get("resolved",0)), ("Vencidos", sm.get("overdue",0)), ("Sem responsável", sm.get("unassigned",0)), ("Erros", sm.get("errors",0))]
    x = 0.45
    for label, value in metrics:
        shape = slide.shapes.add_shape(1, Inches(x), Inches(1.35), Inches(1.95), Inches(0.95))
        shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor(248,250,252); shape.line.color.rgb = RGBColor(226,232,240)
        tb = slide.shapes.add_textbox(Inches(x+0.08), Inches(1.48), Inches(1.75), Inches(0.5))
        p = tb.text_frame.paragraphs[0]; p.text = str(value); p.font.size = Pt(22); p.font.bold = True; p.font.color.rgb = red if label in {"Vencidos", "Erros"} else navy
        tb2 = slide.shapes.add_textbox(Inches(x+0.08), Inches(1.93), Inches(1.75), Inches(0.25))
        q = tb2.text_frame.paragraphs[0]; q.text = label; q.font.size = Pt(9); q.font.bold = True; q.font.color.rgb = gray
        x += 2.05
    y = 2.6
    txt = slide.shapes.add_textbox(Inches(0.55), Inches(y), Inches(6.2), Inches(1.6))
    tf = txt.text_frame; tf.text = "Leitura executiva"
    tf.paragraphs[0].font.bold = True; tf.paragraphs[0].font.size = Pt(14); tf.paragraphs[0].font.color.rgb = navy
    for line in report.get("executive_reading", [])[:5]:
        p = tf.add_paragraph(); p.text = f"• {line}"; p.font.size = Pt(10); p.font.color.rgb = RGBColor(31,41,55)
    txt2 = slide.shapes.add_textbox(Inches(7.05), Inches(y), Inches(5.8), Inches(2.8))
    tf2 = txt2.text_frame; tf2.text = "Principais responsáveis / severidades"
    tf2.paragraphs[0].font.bold = True; tf2.paragraphs[0].font.size = Pt(14); tf2.paragraphs[0].font.color.rgb = navy
    for row in (report.get("by_owner", [])[:5] + report.get("by_severity", [])[:4]):
        label = row.get("owner") or row.get("severity") or "—"
        p = tf2.add_paragraph(); p.text = f"• {label}: {row.get('count',0)}"; p.font.size = Pt(10); p.font.color.rgb = RGBColor(31,41,55)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Fila crítica de SLA", "Alertas vencidos, próximos do vencimento ou sem responsável")
    rows = report.get("critical_queue", [])[:10]
    table = slide.shapes.add_table(len(rows)+1 if rows else 2, 5, Inches(0.45), Inches(1.25), Inches(12.4), Inches(5.4)).table
    headers = ["Data", "SLA", "Severidade", "Responsável", "Título"]
    for i, h in enumerate(headers):
        table.cell(0,i).text = h
    if rows:
        for r, row in enumerate(rows, start=1):
            values = [row.get("created_at_display") or "—", row.get("sla_status") or "—", row.get("severity") or "—", row.get("assigned_to") or "Sem responsável", str(row.get("title") or "—")[:90]]
            for c, v in enumerate(values): table.cell(r,c).text = v
    else:
        for c, v in enumerate(["—", "—", "—", "—", "Nenhum alerta crítico no período"]): table.cell(1,c).text = v
    for row in table.rows:
        for cell in row.cells:
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(8)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
