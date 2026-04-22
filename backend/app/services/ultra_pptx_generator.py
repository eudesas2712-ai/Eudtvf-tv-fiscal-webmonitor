from io import BytesIO

from fastapi.responses import StreamingResponse
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


def build_market_report_pptx(project_id: str, summary: dict) -> BytesIO:
    total_banners = summary.get("total_banners", 0) or 0
    total_investment = summary.get("total_investment", 0) or 0
    share_of_voice = summary.get("share_of_voice", []) or []
    portal_ranking = summary.get("portal_ranking", []) or []
    confidence_summary = summary.get("confidence_summary", []) or []
    strategic_alerts = summary.get("strategic_alerts", []) or []

    advertisers_count = len(share_of_voice)
    portals_count = len(portal_ranking)

    top_advertiser = share_of_voice[0] if share_of_voice else {}
    top_portal = portal_ranking[0] if portal_ranking else {}

    has_data = total_banners > 0 and (advertisers_count > 0 or portals_count > 0)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    RED = RGBColor(197, 38, 37)
    DARK_RED = RGBColor(143, 20, 24)
    DARK = RGBColor(31, 41, 55)
    GRAY = RGBColor(107, 114, 128)
    LIGHT = RGBColor(245, 247, 250)
    WHITE = RGBColor(255, 255, 255)
    BORDER = RGBColor(230, 232, 236)

    def brl(value: float | int) -> str:
        return f"R$ {float(value or 0):,.0f}".replace(",", ".")

    def add_bg(slide, color=WHITE):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        slide.shapes._spTree.remove(shape._element)
        slide.shapes._spTree.insert(2, shape._element)

    def add_text(
        slide,
        left,
        top,
        width,
        height,
        text,
        size=20,
        bold=False,
        color=DARK,
        align=PP_ALIGN.LEFT,
    ):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = str(text)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = "Arial"
        return box

    def add_section_title(slide, title, subtitle=""):
        add_text(
            slide,
            Inches(0.55),
            Inches(0.30),
            Inches(7.5),
            Inches(0.4),
            "TV Fiscal WebMonitor",
            size=12,
            bold=True,
            color=GRAY,
        )
        add_text(
            slide,
            Inches(0.55),
            Inches(0.68),
            Inches(8.8),
            Inches(0.7),
            title,
            size=28,
            bold=True,
            color=DARK,
        )
        if subtitle:
            add_text(
                slide,
                Inches(0.55),
                Inches(1.18),
                Inches(10.8),
                Inches(0.45),
                subtitle,
                size=13,
                color=GRAY,
            )

    def add_card(slide, left, top, width, height, title, value, value_color=DARK):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = BORDER

        add_text(
            slide,
            left + Inches(0.18),
            top + Inches(0.10),
            width - Inches(0.3),
            Inches(0.22),
            title,
            size=11,
            bold=True,
            color=GRAY,
        )
        add_text(
            slide,
            left + Inches(0.18),
            top + Inches(0.38),
            width - Inches(0.3),
            Inches(0.52),
            value,
            size=23,
            bold=True,
            color=value_color,
        )

    def add_info_card(slide, left, top, width, height, label, title, subtitle):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = BORDER

        add_text(
            slide, left + Inches(0.18), top + Inches(0.12),
            width - Inches(0.3), Inches(0.2),
            label, size=11, bold=True, color=GRAY
        )
        add_text(
            slide, left + Inches(0.18), top + Inches(0.38),
            width - Inches(0.3), Inches(0.35),
            title, size=18, bold=True, color=DARK
        )
        add_text(
            slide, left + Inches(0.18), top + Inches(0.76),
            width - Inches(0.3), Inches(0.45),
            subtitle, size=12, color=GRAY
        )

    def add_bullets(slide, title, items, left=0.7, top=1.8, width=11.7, height=4.8):
        add_text(slide, Inches(left), Inches(1.35), Inches(8), Inches(0.35), title, size=18, bold=True, color=DARK_RED)
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = box.text_frame
        tf.clear()
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = str(item)
            p.level = 0
            p.font.size = Pt(18)
            p.font.name = "Arial"
            p.font.color.rgb = DARK
            p.space_after = Pt(10)

    def add_chart_slide(slide, title, categories, values, x, y, cx, cy, series_name="Share (%)"):
        add_text(slide, Inches(0.7), Inches(1.35), Inches(8), Inches(0.35), title, size=18, bold=True, color=DARK_RED)

        chart_data = CategoryChartData()
        chart_data.categories = categories
        chart_data.add_series(series_name, values)

        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.BAR_CLUSTERED,
            Inches(x), Inches(y), Inches(cx), Inches(cy),
            chart_data
        ).chart

        chart.has_legend = False
        chart.value_axis.visible = True
        chart.value_axis.maximum_scale = max(100, max(values) if values else 100)
        chart.value_axis.minimum_scale = 0
        chart.category_axis.reverse_order = True

        series = chart.series[0]
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RED

        return chart

    def add_table_slide(slide, title, columns, rows, left=0.7, top=1.7, width=11.9, height=4.8):
        add_text(slide, Inches(0.7), Inches(1.3), Inches(8), Inches(0.35), title, size=18, bold=True, color=DARK_RED)
        table_shape = slide.shapes.add_table(
            rows=len(rows) + 1,
            cols=len(columns),
            left=Inches(left),
            top=Inches(top),
            width=Inches(width),
            height=Inches(height),
        )
        table = table_shape.table

        for i, col in enumerate(columns):
            cell = table.cell(0, i)
            cell.text = col
            cell.fill.solid()
            cell.fill.fore_color.rgb = RED
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.color.rgb = WHITE
                    run.font.size = Pt(11)

        for r_idx, row in enumerate(rows, start=1):
            for c_idx, value in enumerate(row):
                cell = table.cell(r_idx, c_idx)
                cell.text = str(value)
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(11)
                        run.font.color.rgb = DARK

    # Slide 1 - Capa
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)

    band = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.55), Inches(0.9), Inches(12.1), Inches(2.15)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = RED
    band.line.fill.background()

    add_text(slide, Inches(0.9), Inches(1.15), Inches(5.5), Inches(0.25), "TV Fiscal WebMonitor", size=13, bold=True, color=WHITE)
    add_text(slide, Inches(0.9), Inches(1.55), Inches(9.5), Inches(0.8), "Inteligência de Mercado Publicitário Digital", size=26, bold=True, color=WHITE)
    add_text(
        slide, Inches(0.9), Inches(2.32), Inches(10.6), Inches(0.5),
        "Relatório executivo com leitura de presença, dominância competitiva, qualidade de identificação e oportunidades comerciais.",
        size=13, color=WHITE
    )

    add_text(slide, Inches(0.9), Inches(3.55), Inches(10), Inches(0.3), f"Projeto monitorado: {project_id}", size=15, bold=True)
    add_text(slide, Inches(0.9), Inches(4.05), Inches(10), Inches(0.3), f"Total de anúncios: {total_banners}", size=18)
    add_text(slide, Inches(0.9), Inches(4.45), Inches(10), Inches(0.3), f"Investimento estimado: {brl(total_investment)}", size=18)
    add_text(slide, Inches(0.9), Inches(4.85), Inches(10), Inches(0.3), f"Anunciantes: {advertisers_count} | Portais: {portals_count}", size=18)
    add_text(slide, Inches(0.9), Inches(6.45), Inches(10.5), Inches(0.25), "TV Fiscal — monitoramento, inteligência e qualificação de presença publicitária.", size=12, color=GRAY)

    # Slide 2 - Resumo executivo
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, LIGHT)
    add_section_title(slide, "Resumo Executivo", "Indicadores consolidados do projeto monitorado")

    add_card(slide, Inches(0.6), Inches(1.8), Inches(2.9), Inches(1.15), "Investimento Total", brl(total_investment), RED)
    add_card(slide, Inches(3.7), Inches(1.8), Inches(2.7), Inches(1.15), "Total de Anúncios", str(total_banners))
    add_card(slide, Inches(6.6), Inches(1.8), Inches(2.5), Inches(1.15), "Anunciantes", str(advertisers_count))
    add_card(slide, Inches(9.3), Inches(1.8), Inches(2.5), Inches(1.15), "Portais", str(portals_count))

    add_info_card(
        slide,
        Inches(0.6), Inches(3.25), Inches(5.55), Inches(1.45),
        "Maior anunciante",
        top_advertiser.get("advertiser", "Sem dados"),
        f"Share: {top_advertiser.get('share_percent', 0)}% | Confiança: {top_advertiser.get('confidence', 'n/a')}",
    )
    add_info_card(
        slide,
        Inches(6.25), Inches(3.25), Inches(5.55), Inches(1.45),
        "Portal líder",
        top_portal.get("portal", "Sem dados"),
        f"Receita estimada: {brl(top_portal.get('investment', 0))}",
    )

    insights = strategic_alerts[:] if strategic_alerts else []
    if not insights:
        if has_data:
            insights = ["O projeto apresenta dados suficientes para leitura executiva de presença publicitária."]
        else:
            insights = [
                "Não há evidências publicitárias registradas para o projeto informado nesta base.",
                "O relatório foi gerado em modo informativo, sem ranking de anunciantes ou portais.",
            ]

    add_bullets(slide, "Leitura Executiva", insights[:4], left=0.7, top=5.15, width=11.6, height=1.8)

    # Slide 3 - Share por anunciante
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    add_section_title(slide, "Dominância por Anunciante", "Top players por share e volume de banners")

    adv_items = share_of_voice[:8] if share_of_voice else [{"advertiser": "Sem dados", "share_percent": 0}]
    adv_categories = [item.get("advertiser", "Sem dados") for item in adv_items]
    adv_values = [float(item.get("share_percent", 0) or 0) for item in adv_items]
    add_chart_slide(slide, "Share por anunciante", adv_categories, adv_values, 0.8, 1.8, 6.1, 4.6)

    adv_rows = [
        [
            idx + 1,
            item.get("advertiser", "Sem dados"),
            item.get("banners", 0),
            f"{item.get('share_percent', 0)}%",
            brl(item.get("investment", 0)),
        ]
        for idx, item in enumerate(share_of_voice[:6])
    ] or [[1, "Sem dados", 0, "0%", brl(0)]]

    add_table_slide(
        slide,
        "Top anunciantes",
        ["Posição", "Anunciante", "Banners", "Share", "Investimento"],
        adv_rows,
        left=7.25,
        top=1.8,
        width=5.35,
        height=4.5,
    )

    # Slide 4 - Share por portal
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    add_section_title(slide, "Performance por Portal", "Distribuição do investimento entre publishers monitorados")

    portal_items = portal_ranking[:8] if portal_ranking else [{"portal": "Sem dados", "share_percent": 0}]
    portal_categories = [item.get("portal", "Sem dados") for item in portal_items]
    portal_values = [float(item.get("share_percent", 0) or 0) for item in portal_items]
    add_chart_slide(slide, "Share por portal", portal_categories, portal_values, 0.8, 1.8, 6.1, 4.6)

    portal_rows = [
        [
            idx + 1,
            item.get("portal", "Sem dados"),
            item.get("banners", 0),
            f"{item.get('share_percent', 0)}%",
            brl(item.get("investment", 0)),
        ]
        for idx, item in enumerate(portal_ranking[:6])
    ] or [[1, "Sem dados", 0, "0%", brl(0)]]

    add_table_slide(
        slide,
        "Top portais",
        ["Posição", "Portal", "Banners", "Share", "Investimento"],
        portal_rows,
        left=7.25,
        top=1.8,
        width=5.35,
        height=4.5,
    )

    # Slide 5 - Confiança da identificação
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, LIGHT)
    add_section_title(slide, "Confiança da Identificação", "Distribuição da qualidade do reconhecimento")

    conf_map = {item.get("confidence"): item for item in confidence_summary}
    conf_high = conf_map.get("alta", {"banners": 0, "investment": 0})
    conf_medium = conf_map.get("media", {"banners": 0, "investment": 0})
    conf_low = conf_map.get("baixa", {"banners": 0, "investment": 0})

    add_card(slide, Inches(0.8), Inches(2.0), Inches(3.6), Inches(1.5), "ALTA", f"{conf_high.get('banners', 0)} banners\n{brl(conf_high.get('investment', 0))}", RED)
    add_card(slide, Inches(4.85), Inches(2.0), Inches(3.6), Inches(1.5), "MÉDIA", f"{conf_medium.get('banners', 0)} banners\n{brl(conf_medium.get('investment', 0))}", DARK)
    add_card(slide, Inches(8.9), Inches(2.0), Inches(3.6), Inches(1.5), "BAIXA", f"{conf_low.get('banners', 0)} banners\n{brl(conf_low.get('investment', 0))}", DARK)

    conf_chart_data = CategoryChartData()
    conf_chart_data.categories = ["Alta", "Média", "Baixa"]
    conf_chart_data.add_series(
        "Banners",
        [
            conf_high.get("banners", 0),
            conf_medium.get("banners", 0),
            conf_low.get("banners", 0),
        ],
    )

    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(2.1), Inches(4.05), Inches(9.1), Inches(2.2),
        conf_chart_data
    ).chart
    chart.has_legend = False
    chart.value_axis.minimum_scale = 0
    series = chart.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = RED

    # Slide 6 - Conclusão
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)
    add_section_title(slide, "Conclusão Estratégica", "Síntese final para tomada de decisão")

    if has_data:
        final_points = [
            "O recorte analisado evidencia concentração relevante de inventário e forte peso de registros não qualificados.",
            "Há oportunidade clara de ampliar a identificação de anunciantes para elevar a inteligência competitiva.",
            "Os rankings de share e os portais líderes indicam caminhos práticos para priorização de canais e otimização de presença.",
        ]
    else:
        final_points = [
            "Não há evidências publicitárias registradas para o projeto informado nesta base.",
            "A apresentação foi gerada em modo informativo para preservar a disponibilidade do endpoint.",
            "Recomendação: validar a ingestão de banners, restaurar a base histórica ou reprocessar as coletas antes da análise executiva.",
        ]

    add_bullets(slide, "Conclusão", final_points, left=0.8, top=1.8, width=11.2, height=3.4)
    add_text(
        slide,
        Inches(0.8), Inches(6.45), Inches(11), Inches(0.25),
        "Relatório gerado automaticamente pela plataforma TV Fiscal WebMonitor.",
        size=12, color=GRAY
    )

    bio = BytesIO()
    prs.save(bio)
    bio.seek(0)
    return bio


def build_market_report_pptx_response(project_id: str, summary: dict) -> StreamingResponse:
    bio = build_market_report_pptx(project_id, summary)
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={
            "Content-Disposition": f'attachment; filename="intel_report_{project_id}.pptx"'
        },
    )
