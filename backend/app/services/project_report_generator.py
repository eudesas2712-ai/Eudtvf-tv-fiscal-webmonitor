from __future__ import annotations

import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any

from fastapi.responses import FileResponse
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, PageBreak
from app.services.visual_report_theme import (
    header_footer as _premium_footer, hero as _premium_hero, kpi_cards as _premium_kpis,
    section_title as _premium_section, notice_box as _premium_notice, bar_list as _premium_bar_list,
    text_panel as _premium_text_panel, table_style as _premium_table_style, short as _p_short,
    brl as _p_brl, pct as _p_pct, styles as _premium_styles
)

RED = RGBColor(197, 38, 37)
DARK = RGBColor(31, 41, 55)
GRAY = RGBColor(105, 112, 120)
LIGHT = RGBColor(244, 246, 248)
WHITE = RGBColor(255, 255, 255)
MID = RGBColor(226, 230, 235)


def _brl(value: Any) -> str:
    try:
        return f"R$ {float(value or 0):,.0f}".replace(",", ".")
    except Exception:
        return "R$ 0"


def _pct(value: Any) -> str:
    try:
        return f"{float(value or 0):.1f}%"
    except Exception:
        return "0,0%"


def _short(text: Any, size: int = 42) -> str:
    value = str(text or "—")
    return value if len(value) <= size else value[: size - 1] + "…"


def _filename(prefix: str, project_id: str, ext: str) -> str:
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    safe_project = str(project_id).replace("/", "-")[:36]
    return f"{prefix}_{safe_project}_{stamp}.{ext}"


def _base_table_style() -> TableStyle:
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#b00020")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
        ("BACKGROUND", (0, 1), (-1, -1), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f9fafb")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ])


def build_project_pdf(data: dict, output_path: str) -> str:
    """Premium unified WebMonitor PDF.

    V21 aligns the project/publicity/checking report with the Editorial Premium
    Dashboard visual system: navy hero, transparent logo, KPI strip, executive
    panels, modern bar lists, and cleaner evidence tables.
    """
    doc = SimpleDocTemplate(
        output_path,
        pagesize=landscape(A4),
        rightMargin=14 * mm,
        leftMargin=14 * mm,
        topMargin=10 * mm,
        bottomMargin=14 * mm,
    )
    story: list[Any] = []
    st = _premium_styles()

    project = data.get("project", {}) or {}
    totals = data.get("totals", {}) or {}
    quality = data.get("data_quality", {}) or {}
    insights = data.get("insights", []) or []
    top_advertisers = data.get("top_advertisers", []) or []
    id_quality = data.get("identification_quality", {}) or {}
    top_portals = data.get("top_portals", []) or []
    auditables = data.get("auditables", []) or []
    news_candidates = data.get("news_candidates", []) or []
    editorial = data.get("editorial", {}) or {}
    competitors = data.get("competitors", []) or []
    portals = data.get("linked_portals", []) or []

    project_label = project.get("name") or project.get("id") or "Projeto TV Fiscal WebMonitor"
    subtitle = "Publicidade digital, checking, inteligência de mercado, evidências auditáveis e clipping editorial integrado"
    story.append(_premium_hero(
        "TV Fiscal Intelligence Report",
        subtitle,
        project_label=f"{project_label} · {project.get('client_name') or 'Cliente não definido'}",
        kicker="RELATÓRIO COMPLETO · TV FISCAL WEBMONITOR",
    ))
    story.append(Spacer(1, 4 * mm))
    story.append(_premium_kpis([
        ("Detectados", totals.get("detected", 0), "itens visuais"),
        ("Mercado", totals.get("market_items", 0), "inteligência"),
        ("Auditáveis", totals.get("auditables", 0), "checking"),
        ("Notícias cand.", totals.get("news_candidates", 0), "editorial"),
        ("Matérias", totals.get("editorial_items", 0), "clipping"),
        ("Investimento", _p_brl(totals.get("investment", 0)), "estimado"),
    ]))
    story.append(Spacer(1, 4 * mm))

    for warning in (quality.get("warnings") or [])[:2]:
        story.append(_premium_notice(warning, tone="warning"))
        story.append(Spacer(1, 2 * mm))
    if id_quality.get("unknown_items"):
        story.append(_premium_notice(
            f"{id_quality.get('unknown_items')} item(ns) permanecem pendente(s) de identificação comercial; o bloco não é tratado como anunciante líder.",
            tone="warning",
        ))
        story.append(Spacer(1, 2 * mm))

    story.append(Table([[
        _premium_bar_list("Top anunciantes identificados", top_advertisers, "name", "items", sub_key="investment_label", limit=5),
        _premium_bar_list("Top portais", top_portals, "portal", "items", sub_key="investment_label", limit=5),
        _premium_text_panel("Leitura executiva", insights[:6] or ["Sem leitura executiva disponível para o recorte atual."], width=64 * mm),
    ]], colWidths=[94 * mm, 94 * mm, 66 * mm], rowHeights=[68 * mm], style=TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ])))
    story.append(Spacer(1, 5 * mm))

    story.append(Table([[
        _premium_text_panel("Escopo operacional", [
            f"Modo de análise: {quality.get('label') or 'Mercado amplo'}",
            f"Cliente: {project.get('client_name') or 'não definido'}",
            f"Segmento: {project.get('segment_name') or 'não definido'}",
            f"Portais vinculados: {len(portals)}",
        ], width=124 * mm),
        _premium_text_panel("Editorial integrado", [
            f"Matérias coletadas: {editorial.get('total_items', 0)}",
            f"Com termos/marcas: {editorial.get('matched_items', 0)}",
            f"Fontes editoriais: {editorial.get('sources_count', 0)}",
            f"Temas recorrentes: {editorial.get('topics_count', 0)}",
        ], width=124 * mm),
    ]], colWidths=[127 * mm, 127 * mm], style=TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ])))

    story.append(PageBreak())
    story.append(_premium_section("Configuração operacional do projeto"))
    cfg_rows = [["Portais vinculados", "Anunciantes vinculados / papel"]]
    portal_text = ", ".join([p.get("name") or p.get("base_url") or "—" for p in portals[:12]]) or "Nenhum portal vinculado"
    adv_text = ", ".join([f"{a.get('name')} ({a.get('role') or 'monitorado'})" for a in competitors[:14]]) or "Nenhum anunciante vinculado"
    cfg_rows.append([_p_short(portal_text, 150), _p_short(adv_text, 150)])
    table = Table(cfg_rows, colWidths=[125 * mm, 125 * mm])
    table.setStyle(_premium_table_style())
    story.append(table)

    story.append(_premium_section("Ranking de anunciantes identificados — inteligência de mercado"))
    rows = [["#", "Anunciante", "Papel", "Evid.", "Audit.", "Invest.", "Share", "Portais"]]
    for idx, item in enumerate(top_advertisers[:14], start=1):
        rows.append([
            idx,
            _p_short(item.get("name"), 34),
            item.get("role") or "—",
            item.get("items", 0),
            item.get("auditables", 0),
            _p_brl(item.get("investment", 0)),
            _p_pct(item.get("share", 0)),
            item.get("portals", 0),
        ])
    if len(rows) == 1:
        rows.append(["—", "Sem dados", "—", 0, 0, _p_brl(0), _p_pct(0), 0])
    table = Table(rows, colWidths=[10 * mm, 62 * mm, 34 * mm, 18 * mm, 18 * mm, 30 * mm, 22 * mm, 20 * mm])
    table.setStyle(_premium_table_style())
    story.append(table)

    story.append(_premium_section("Portais com presença detectada"))
    rows = [["#", "Portal", "Evid.", "Audit.", "Invest.", "Share"]]
    for idx, item in enumerate(top_portals[:12], start=1):
        rows.append([idx, _p_short(item.get("portal"), 72), item.get("items", 0), item.get("auditables", 0), _p_brl(item.get("investment", 0)), _p_pct(item.get("share", 0))])
    if len(rows) == 1:
        rows.append(["—", "Sem dados", 0, 0, _p_brl(0), _p_pct(0)])
    table = Table(rows, colWidths=[10 * mm, 105 * mm, 22 * mm, 22 * mm, 34 * mm, 24 * mm])
    table.setStyle(_premium_table_style())
    story.append(table)

    story.append(_premium_section("Evidências auditáveis para checking"))
    rows = [["Data", "Anunciante", "Portal", "Formato", "Valor", "Evidência"]]
    for item in auditables[:14]:
        rows.append([
            _p_short(item.get("created_at"), 18),
            _p_short(item.get("advertiser"), 34),
            _p_short(item.get("portal"), 42),
            item.get("format", "—"),
            _p_brl(item.get("estimated_value", 0)),
            item.get("evidence_type") or "preservada",
        ])
    if len(rows) == 1:
        rows.append(["—", "Sem evidências auditáveis", "—", "—", _p_brl(0), "—"])
    table = Table(rows, colWidths=[28 * mm, 55 * mm, 55 * mm, 25 * mm, 28 * mm, 35 * mm])
    table.setStyle(_premium_table_style())
    story.append(table)

    story.append(PageBreak())
    story.append(_premium_section("Candidatos a monitoramento de notícias"))
    rows = [["Data", "Título/OCR", "Portal", "Score notícia", "Página"]]
    for item in news_candidates[:12]:
        rows.append([
            _p_short(item.get("created_at"), 18),
            _p_short(item.get("title"), 78),
            _p_short(item.get("portal"), 44),
            item.get("news_score", 0),
            _p_short(item.get("page_url"), 55),
        ])
    if len(rows) == 1:
        rows.append(["—", "Sem candidatos a notícia", "—", 0, "—"])
    table = Table(rows, colWidths=[28 * mm, 85 * mm, 48 * mm, 28 * mm, 70 * mm])
    table.setStyle(_premium_table_style())
    story.append(table)

    story.append(_premium_section("Resumo editorial / clipping"))
    sentiment_counts = editorial.get("sentiment_counts", {}) or {}
    sent_label = ", ".join([f"{k}: {v}" for k, v in sentiment_counts.items()]) or "Sem dados de sentimento"
    editorial_rows = [["Matérias", "Com termos", "Fontes", "Temas", "Score médio", "Sentimento"]]
    editorial_rows.append([
        editorial.get("total_items", 0),
        editorial.get("matched_items", 0),
        editorial.get("sources_count", 0),
        editorial.get("topics_count", 0),
        editorial.get("avg_editorial_score", 0),
        _p_short(sent_label, 85),
    ])
    table = Table(editorial_rows, colWidths=[25 * mm, 25 * mm, 22 * mm, 22 * mm, 26 * mm, 100 * mm])
    table.setStyle(_premium_table_style())
    story.append(table)

    story.append(_premium_section("Temas e termos editoriais"))
    rows = [["Tema", "Matérias", "Termo/Marca", "Matérias"]]
    topics = editorial.get("top_topics", []) or []
    terms = editorial.get("top_terms", []) or []
    max_rows = max(len(topics[:8]), len(terms[:8]), 1)
    for i in range(max_rows):
        topic = topics[i] if i < len(topics) else {}
        term = terms[i] if i < len(terms) else {}
        rows.append([_p_short(topic.get("topic") or "—", 40), topic.get("items", "—"), _p_short(term.get("term") or "—", 40), term.get("items", "—")])
    table = Table(rows, colWidths=[75 * mm, 22 * mm, 75 * mm, 22 * mm])
    table.setStyle(_premium_table_style())
    story.append(table)

    story.append(_premium_section("Últimas matérias monitoradas"))
    rows = [["Data", "Fonte", "Tema", "Sent.", "Termos", "Título"]]
    for item in (editorial.get("latest_items", []) or [])[:12]:
        rows.append([
            _p_short(item.get("created_at"), 18),
            _p_short(item.get("source_name"), 30),
            _p_short(item.get("topic"), 22),
            _p_short(item.get("sentiment"), 16),
            _p_short(item.get("terms_label"), 35),
            _p_short(item.get("title"), 72),
        ])
    if len(rows) == 1:
        rows.append(["—", "Sem matérias editoriais", "—", "—", "—", "—"])
    table = Table(rows, colWidths=[28 * mm, 42 * mm, 26 * mm, 22 * mm, 42 * mm, 82 * mm])
    table.setStyle(_premium_table_style())
    story.append(table)

    doc.build(story, onFirstPage=_premium_footer, onLaterPages=_premium_footer)
    return output_path

def _add_bg(slide, color=WHITE):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)


def _add_text(slide, left, top, width, height, text, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Arial"
    return box


def _add_title(slide, title, subtitle=None):
    _add_text(slide, Inches(0.65), Inches(0.58), Inches(12), Inches(0.45), title, 26, True, DARK)
    if subtitle:
        _add_text(slide, Inches(0.68), Inches(1.08), Inches(11.7), Inches(0.35), subtitle, 12, False, GRAY)


def _add_card(slide, left, top, width, height, label, value):
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = MID
    _add_text(slide, left + Inches(0.18), top + Inches(0.13), width - Inches(0.3), Inches(0.25), label, 9, True, GRAY)
    _add_text(slide, left + Inches(0.18), top + Inches(0.43), width - Inches(0.3), Inches(0.5), value, 19, True, RED)


def _add_table(slide, title, headers, rows, left=Inches(0.65), top=Inches(1.55), width=Inches(12.0), height=Inches(4.9)):
    _add_text(slide, left, top - Inches(0.35), width, Inches(0.3), title, 15, True, RED)
    rows = rows[:8] if rows else [["—"] * len(headers)]
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), left, top, width, height)
    table = table_shape.table
    for c, header in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = str(header)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RED
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(8)
                r.font.bold = True
                r.font.color.rgb = WHITE
                r.font.name = "Arial"
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = str(value)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r_idx % 2 else LIGHT
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
                    r.font.color.rgb = DARK
                    r.font.name = "Arial"


def build_project_pptx(data: dict, output_path: str) -> str:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    project = data.get("project", {}) or {}
    totals = data.get("totals", {}) or {}
    quality = data.get("data_quality", {}) or {}
    insights = data.get("insights", []) or []
    top_advertisers = data.get("top_advertisers", []) or []
    id_quality = data.get("identification_quality", {}) or {}
    top_portals = data.get("top_portals", []) or []
    auditables = data.get("auditables", []) or []
    news_candidates = data.get("news_candidates", []) or []
    editorial = data.get("editorial", {}) or {}
    competitors = data.get("competitors", []) or []
    portals = data.get("linked_portals", []) or []

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, LIGHT)
    hero = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.25), Inches(12.0), Inches(2.25))
    hero.fill.solid()
    hero.fill.fore_color.rgb = RED
    hero.line.fill.background()
    _add_text(slide, Inches(1.05), Inches(1.72), Inches(10.8), Inches(0.42), "TV Fiscal WebMonitor", 18, True, WHITE)
    _add_text(slide, Inches(1.05), Inches(2.15), Inches(10.8), Inches(0.62), "Relatório Completo do Projeto", 31, True, WHITE)
    _add_text(slide, Inches(1.05), Inches(2.9), Inches(10.8), Inches(0.32), _short(project.get("name"), 90), 13, False, WHITE)
    _add_text(slide, Inches(0.9), Inches(4.25), Inches(11.8), Inches(0.36), f"Cliente: {project.get('client_name') or 'Não definido'} · Segmento: {project.get('segment_name') or 'Não definido'}", 14, False, DARK)
    _add_text(slide, Inches(0.9), Inches(4.62), Inches(11.8), Inches(0.36), f"Modo de análise: {quality.get('label') or 'Mercado amplo'}", 13, True, RED)
    _add_text(slide, Inches(0.9), Inches(5.0), Inches(11.8), Inches(0.36), f"Investimento: {_brl(totals.get('investment'))} · Auditáveis: {totals.get('auditables', 0)} · Matérias: {totals.get('editorial_items', 0)}", 16, True, DARK)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, WHITE)
    _add_title(slide, "Resumo executivo", "Indicadores consolidados do projeto")
    _add_card(slide, Inches(0.7), Inches(1.65), Inches(1.8), Inches(1.05), "Detectados", str(totals.get("detected", 0)))
    _add_card(slide, Inches(2.7), Inches(1.65), Inches(1.8), Inches(1.05), "Analisados", str(totals.get("market_items", 0)))
    _add_card(slide, Inches(4.7), Inches(1.65), Inches(1.8), Inches(1.05), "Auditáveis", str(totals.get("auditables", 0)))
    _add_card(slide, Inches(6.7), Inches(1.65), Inches(1.8), Inches(1.05), "Matérias", str(totals.get("editorial_items", 0)))
    _add_card(slide, Inches(8.7), Inches(1.65), Inches(1.8), Inches(1.05), "Termos", str(totals.get("editorial_matched_items", 0)))
    _add_card(slide, Inches(10.7), Inches(1.65), Inches(2.0), Inches(1.05), "Invest.", _brl(totals.get("investment", 0)))
    y = Inches(3.1)
    if quality.get("warnings"):
        _add_text(slide, Inches(0.85), y, Inches(11.8), Inches(0.5), f"Atenção: {quality.get('warnings', [''])[0]}", 11.5, True, RED)
        y += Inches(0.62)
    _add_text(slide, Inches(0.85), y - Inches(0.1), Inches(11.8), Inches(0.3), "Leitura executiva", 16, True, RED)
    y += Inches(0.36)
    for item in (insights or ["Sem insights disponíveis."])[:6]:
        _add_text(slide, Inches(0.95), y, Inches(11.5), Inches(0.35), f"• {item}", 12.0, False, DARK)
        y += Inches(0.40)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, LIGHT)
    _add_title(slide, "Configuração operacional", "Portais e anunciantes vinculados ao projeto")
    _add_table(
        slide,
        "Portais vinculados",
        ["Portal", "URL"],
        [[_short(p.get("name"), 36), _short(p.get("base_url"), 80)] for p in portals[:8]],
        left=Inches(0.65), top=Inches(1.55), width=Inches(6.0), height=Inches(4.8),
    )
    _add_table(
        slide,
        "Anunciantes vinculados",
        ["Anunciante", "Papel", "Segmento"],
        [[_short(a.get("name"), 28), a.get("role") or "monitorado", _short(a.get("segment_name"), 22)] for a in competitors[:8]],
        left=Inches(6.85), top=Inches(1.55), width=Inches(5.8), height=Inches(4.8),
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, WHITE)
    _add_title(slide, "Inteligência de mercado", "Ranking por anunciante")
    _add_table(slide, "Top anunciantes identificados", ["#", "Anunciante", "Evid.", "Audit.", "Invest.", "Share"], [[idx + 1, _short(a.get("name"), 28), a.get("items", 0), a.get("auditables", 0), _brl(a.get("investment", 0)), _pct(a.get("share", 0))] for idx, a in enumerate(top_advertisers[:8])])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, LIGHT)
    _add_title(slide, "Presença por portal", "Distribuição das evidências de mercado")
    _add_table(slide, "Top portais", ["#", "Portal", "Evid.", "Audit.", "Invest.", "Share"], [[idx + 1, _short(p.get("portal"), 42), p.get("items", 0), p.get("auditables", 0), _brl(p.get("investment", 0)), _pct(p.get("share", 0))] for idx, p in enumerate(top_portals[:8])])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, WHITE)
    _add_title(slide, "Evidências auditáveis", "Itens prontos para checking publicitário")
    _add_table(slide, "Últimas evidências auditáveis", ["Data", "Anunciante", "Portal", "Formato", "Valor"], [[_short(e.get("created_at"), 18), _short(e.get("advertiser"), 28), _short(e.get("portal"), 34), e.get("format", "—"), _brl(e.get("estimated_value", 0))] for e in auditables[:8]])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, LIGHT)
    _add_title(slide, "Candidatos a notícia", "Itens editoriais preservados para o módulo de monitoramento de notícias")
    _add_table(slide, "Candidatos editoriais", ["Data", "Título/OCR", "Portal", "Score"], [[_short(n.get("created_at"), 18), _short(n.get("title"), 62), _short(n.get("portal"), 32), n.get("news_score", 0)] for n in news_candidates[:8]])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, WHITE)
    _add_title(slide, "Resumo editorial", "Clipping, termos monitorados, sentimentos e temas")
    _add_card(slide, Inches(0.7), Inches(1.65), Inches(2.0), Inches(1.05), "Matérias", str(editorial.get("total_items", 0)))
    _add_card(slide, Inches(2.95), Inches(1.65), Inches(2.0), Inches(1.05), "Com termos", str(editorial.get("matched_items", 0)))
    _add_card(slide, Inches(5.2), Inches(1.65), Inches(2.0), Inches(1.05), "Fontes", str(editorial.get("sources_count", 0)))
    _add_card(slide, Inches(7.45), Inches(1.65), Inches(2.0), Inches(1.05), "Temas", str(editorial.get("topics_count", 0)))
    _add_card(slide, Inches(9.7), Inches(1.65), Inches(2.8), Inches(1.05), "Score médio", str(editorial.get("avg_editorial_score", 0)))
    _add_table(
        slide,
        "Temas recorrentes",
        ["Tema", "Matérias"],
        [[_short(t.get("topic"), 52), t.get("items", 0)] for t in (editorial.get("top_topics", []) or [])[:8]],
        left=Inches(0.8), top=Inches(3.25), width=Inches(5.6), height=Inches(3.2),
    )
    _add_table(
        slide,
        "Termos/marcas detectados",
        ["Termo", "Matérias"],
        [[_short(t.get("term"), 52), t.get("items", 0)] for t in (editorial.get("top_terms", []) or [])[:8]],
        left=Inches(6.8), top=Inches(3.25), width=Inches(5.7), height=Inches(3.2),
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, LIGHT)
    _add_title(slide, "Últimas matérias monitoradas", "Resumo executivo do clipping editorial")
    _add_table(
        slide,
        "Matérias recentes",
        ["Data", "Fonte", "Tema", "Sent.", "Termos", "Título"],
        [[
            _short(i.get("created_at"), 16),
            _short(i.get("source_name"), 24),
            _short(i.get("topic"), 18),
            _short(i.get("sentiment"), 12),
            _short(i.get("terms_label"), 26),
            _short(i.get("title"), 52),
        ] for i in (editorial.get("latest_items", []) or [])[:8]],
        left=Inches(0.55), top=Inches(1.55), width=Inches(12.3), height=Inches(4.95),
    )

    prs.save(output_path)
    return output_path


def build_project_pdf_response(data: dict) -> FileResponse:
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
    tmp.close()
    build_project_pdf(data, tmp.name)
    return FileResponse(tmp.name, media_type="application/pdf", filename=_filename("relatorio_completo_projeto", data.get("project", {}).get("id") or "projeto", "pdf"))


def build_project_pptx_response(data: dict) -> FileResponse:
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pptx")
    tmp.close()
    build_project_pptx(data, tmp.name)
    return FileResponse(tmp.name, media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation", filename=_filename("relatorio_completo_projeto", data.get("project", {}).get("id") or "projeto", "pptx"))
